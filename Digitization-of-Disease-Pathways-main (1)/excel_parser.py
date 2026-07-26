# excel_parser.py

import pandas as pd
import openpyxl
from typing import Dict, List, Tuple, Any, Optional
from collections import OrderedDict

class CADExcelParser:
    def __init__(self, excel_path: str, sheet_name: str = None):
        """
        Initialize Excel parser with dynamic sheet selection
        Args:
            excel_path: Path to the Excel file
            sheet_name: Optional sheet name. If None, uses active sheet.
                       If provided but not found, falls back to active sheet with warning.
        """
        self.workbook = openpyxl.load_workbook(excel_path)
        
        # Dynamic sheet selection
        if sheet_name is None:
            self.worksheet = self.workbook.active
            print(f"Using active sheet: {self.worksheet.title}")
        else:
            if sheet_name in self.workbook.sheetnames:
                self.worksheet = self.workbook[sheet_name]
                print(f"Using specified sheet: {sheet_name}")
            else:
                self.worksheet = self.workbook.active
                print(f"Warning: Sheet '{sheet_name}' not found. Available sheets: {self.workbook.sheetnames}")
                print(f"Using active sheet: {self.worksheet.title}")
        
        # Updated column mapping for new structure
        self.solution_columns = {
            'E': 'digitalization',
            'F': 'automation', 
            'G': 'sensing',
            'H': 'clinical_innovation',
            'I': 'process_innovation'
        }
    
    def parse_disease_pathway(self) -> Dict[str, Any]:
        """Parse Excel with new structure, preserving stage order"""
        stage_ranges = self._get_stage_ranges()
        parsed_data = {
            'disease_name': '',
            'sections': OrderedDict()  # Use OrderedDict to preserve order
        }
        
        # Process stages in the order they appear in Excel
        for stage_name, (start_row, end_row) in stage_ranges.items():
            stage_data = self._parse_stage_data(stage_name, start_row, end_row)
            parsed_data['sections'][stage_name.lower()] = stage_data
            
        return parsed_data
    
    def _get_stage_ranges(self) -> OrderedDict:
        """Get stage ranges from merged cells in column A, preserving Excel row order"""
        stage_list = []
        
        # Collect all stage ranges with their starting row numbers
        for merged_range in self.worksheet.merged_cells.ranges:
            if merged_range.min_col == 1 and merged_range.max_col == 1:  # Column A only
                stage_cell = self.worksheet.cell(merged_range.min_row, 1)
                if stage_cell.value and str(stage_cell.value).strip():
                    stage_name = str(stage_cell.value).strip()
                    # Skip header rows and "Stage" title
                    if stage_name.lower() not in ['stage'] and merged_range.min_row > 2:
                        stage_list.append((
                            merged_range.min_row,  # Row number for sorting
                            stage_name,
                            (merged_range.min_row, merged_range.max_row)
                        ))
        
        # Sort by row number to preserve Excel top-to-bottom order
        stage_list.sort(key=lambda x: x[0])
        
        # Convert to OrderedDict to maintain order
        stage_ranges = OrderedDict()
        for _, stage_name, range_tuple in stage_list:
            stage_ranges[stage_name] = range_tuple
            
        print(f"📊 Found stages in order: {list(stage_ranges.keys())}")
        return stage_ranges
    
    def _parse_stage_data(self, stage_name: str, start_row: int, end_row: int) -> Dict[str, Any]:
        """Parse stage data with new fields including J, K, L columns"""
        stage_data = {
            'name': stage_name,
            'data': [],
            'row_count': 0,
            'overview': self._get_merged_cell_value(2, start_row, end_row),  # Column B
            'stakeholders': self._get_merged_cell_value(3, start_row, end_row)  # Column C
        }
        
        for row in range(start_row, end_row + 1):
            row_data = self._parse_row_data(row)
            if row_data.get('pain_point'):
                stage_data['data'].append(row_data)
                stage_data['row_count'] += 1
                
        return stage_data
    
    def _get_merged_cell_value(self, col_num: int, start_row: int, end_row: int) -> str:
        """Get value from merged cell"""
        for merged_range in self.worksheet.merged_cells.ranges:
            if (merged_range.min_col == col_num and merged_range.max_col == col_num and
                merged_range.min_row <= start_row and merged_range.max_row >= end_row):
                cell = self.worksheet.cell(merged_range.min_row, col_num)
                return str(cell.value).strip() if cell.value else ""
        
        cell = self.worksheet.cell(start_row, col_num)
        return str(cell.value).strip() if cell.value else ""
    
    def _parse_row_data(self, row: int) -> Dict[str, Any]:
        """Parse individual row data including new J, K, L columns"""
        row_data = {
            'row_number': row,
            'pain_point': self._get_cell_value(row, 4),  # Column D
            'solutions': {},
            # NEW: Separate fields for J, K, L columns
            'coverage': self._get_cell_value(row, 10),  # Column J - Our portfolio coverage
            'existing_solutions': self._get_cell_value(row, 11),  # Column K - Ecosystem solutions
            'sources': self._get_cell_value(row, 12)  # Column L - Sources
        }
        
        # Get solutions from columns E-I (kept separate from J, K, L)
        for col_letter, solution_type in self.solution_columns.items():
            col_num = ord(col_letter) - ord('A') + 1
            row_data['solutions'][solution_type] = self._get_cell_value(row, col_num)
        
        return row_data
    
    def _get_cell_value(self, row: int, col: int) -> str:
        """Get cell value as string"""
        cell = self.worksheet.cell(row, col)
        return str(cell.value).strip() if cell.value and str(cell.value).strip() != "None" else ""
    
    def get_summary_stats(self, parsed_data: Dict[str, Any]) -> Dict[str, Any]:
        """Get summary statistics including new fields"""
        stats = {
            'total_stages': len(parsed_data['sections']),
            'total_pain_points': 0,
            'total_solutions': 0,
            'total_coverage_entries': 0,  # NEW: Count coverage entries
            'total_ecosystem_solutions': 0,  # NEW: Count ecosystem solutions
            'total_sources': 0,  # NEW: Count sources
            'stages': []
        }
        
        for stage_name, stage_data in parsed_data['sections'].items():
            stage_stats = {
                'name': stage_name,
                'pain_points_count': stage_data['row_count'],
                'solutions_count': 0,
                'coverage_count': 0,  # NEW
                'ecosystem_solutions_count': 0,  # NEW
                'sources_count': 0  # NEW
            }
            
            for row_data in stage_data['data']:
                # Count solutions (E-I columns)
                for solution_text in row_data['solutions'].values():
                    if solution_text:
                        stage_stats['solutions_count'] += 1
                
                # Count new fields (J, K, L columns)
                if row_data.get('coverage'):
                    stage_stats['coverage_count'] += 1
                    
                if row_data.get('existing_solutions'):
                    stage_stats['ecosystem_solutions_count'] += 1
                    
                if row_data.get('sources'):
                    stage_stats['sources_count'] += 1
            
            stats['stages'].append(stage_stats)
            stats['total_pain_points'] += stage_stats['pain_points_count']
            stats['total_solutions'] += stage_stats['solutions_count']
            stats['total_coverage_entries'] += stage_stats['coverage_count']
            stats['total_ecosystem_solutions'] += stage_stats['ecosystem_solutions_count']
            stats['total_sources'] += stage_stats['sources_count']
        
        return stats
    
    def get_available_sheets(self) -> List[str]:
        """Get list of all available sheet names in the workbook"""
        return self.workbook.sheetnames
    
    def get_current_sheet_name(self) -> str:
        """Get currently selected sheet name"""
        return self.worksheet.title
    
    def get_parsed_data_summary(self, parsed_data: Dict[str, Any]) -> str:
        """Get a formatted summary of parsed data including new fields"""
        summary_lines = []
        summary_lines.append(f"Disease: {parsed_data['disease_name'].upper()}")
        summary_lines.append(f"Total Stages: {len(parsed_data['sections'])}")
        summary_lines.append("="*50)
        
        for stage_name, stage_data in parsed_data['sections'].items():
            summary_lines.append(f"\n🏥 STAGE: {stage_name.upper()}")
            summary_lines.append(f"   Pain Points: {stage_data['row_count']}")
            summary_lines.append(f"   Overview: {stage_data['overview'][:100]}..." if len(stage_data['overview']) > 100 else f"   Overview: {stage_data['overview']}")
            
            # Show sample data with new fields
            if stage_data['data']:
                sample = stage_data['data'][0]
                summary_lines.append(f"   Sample Pain Point: {sample['pain_point'][:80]}...")
                
                # Show new fields
                if sample.get('coverage'):
                    summary_lines.append(f"   Portfolio Coverage: {sample['coverage'][:60]}...")
                if sample.get('existing_solutions'):
                    summary_lines.append(f"   Ecosystem Solutions: {sample['existing_solutions'][:60]}...")
                if sample.get('sources'):
                    summary_lines.append(f"   Sources: {sample['sources'][:60]}...")
        
        return "\n".join(summary_lines)
